import sys
from pathlib import Path
from datetime import datetime

from openpyxl import Workbook, load_workbook

import database


def get_base_directory():
    """
    Return the directory where persistent application data should live.
    """

    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent

    return Path(__file__).resolve().parent


BASE_DIR = get_base_directory()

EXCEL_FILE = BASE_DIR / "DesktopTimeTracker.xlsx"


def create_excel_file():
    """
    Create the Excel report if it does not exist.
    """

    if EXCEL_FILE.exists():
        return

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Desktop Time"

    worksheet["A1"] = "Date"
    worksheet["B1"] = "Desktop"
    worksheet["C1"] = "Total Time"

    worksheet.column_dimensions["A"].width = 15
    worksheet.column_dimensions["B"].width = 40
    worksheet.column_dimensions["C"].width = 18

    workbook.save(EXCEL_FILE)

    workbook.close()


def write_database_to_excel():
    """
    Generate/update the Excel report entirely from SQLite.

    SQLite is the single source of truth.
    """

    database.create_database()
    create_excel_file()

    rows = database.get_all_totals()

    workbook = load_workbook(EXCEL_FILE)

    worksheet = workbook["Desktop Time"]

    # Clear existing data while preserving the header.
    if worksheet.max_row > 1:

        worksheet.delete_rows(
            2,
            worksheet.max_row - 1
        )

    for row_index, (
        date_value,
        desktop_name,
        total_seconds
    ) in enumerate(rows, start=2):

        worksheet.cell(
            row=row_index,
            column=1
        ).value = date_value

        worksheet.cell(
            row=row_index,
            column=2
        ).value = desktop_name

        worksheet.cell(
            row=row_index,
            column=3
        ).value = total_seconds / 86400

        worksheet.cell(
            row=row_index,
            column=3
        ).number_format = "[h]:mm:ss"

    workbook.save(EXCEL_FILE)

    workbook.close()


if __name__ == "__main__":

    database.create_database()

    create_excel_file()

    write_database_to_excel()

    print("Excel report ready:")
    print(EXCEL_FILE)