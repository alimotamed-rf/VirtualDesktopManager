"""
Task Logger
Excel Exporter
"""

import os
import sys

from openpyxl import Workbook
from openpyxl import load_workbook


class ExcelExporter:

    def __init__(self):

        self.file = self.get_excel_path()

    def get_base_path(self):

        if getattr(sys, "frozen", False):

            return os.path.dirname(
                sys.executable
            )

        return os.path.abspath(

            os.path.join(

                os.path.dirname(__file__),

                ".."

            )

        )

    def get_excel_path(self):

        data_folder = os.path.join(

            self.get_base_path(),

            "data"

        )

        os.makedirs(

            data_folder,

            exist_ok=True

        )

        return os.path.join(

            data_folder,

            "WorkLog.xlsx"

        )

    def export(self, projects):

        if os.path.exists(self.file):

            wb = load_workbook(self.file)

            ws = wb.active

        else:

            wb = Workbook()

            ws = wb.active

            ws.append(

                [

                    "Date",

                    "Root",

                    "Project",

                    "Seconds",

                    "Duration"

                ]

            )

        existing = {}

        for row in range(2, ws.max_row + 1):

            key = (

                ws.cell(row, 1).value,

                ws.cell(row, 2).value,

                ws.cell(row, 3).value

            )

            existing[key] = row

        for project in projects.values():

            seconds = int(project["seconds"])

            hours = seconds // 3600

            minutes = (seconds % 3600) // 60

            secs = seconds % 60

            duration = f"{hours:02}:{minutes:02}:{secs:02}"

            from datetime import datetime

            date = datetime.now().strftime("%Y-%m-%d")

            key = (

                date,

                project["root"],

                project["project"]

            )

            if key in existing:

                r = existing[key]

                ws.cell(r, 4).value = seconds

                ws.cell(r, 5).value = duration

            else:

                ws.append(

                    [

                        date,

                        project["root"],

                        project["project"],

                        seconds,

                        duration

                    ]

                )

        try:

            wb.save(self.file)

        except PermissionError:

            pass